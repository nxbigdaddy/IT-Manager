"""Batch UI: multi-row selection, payload shape, idempotent retry and export."""
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from unittest.mock import patch

import pytest
from PySide6.QtCore import QItemSelectionModel, Qt
from PySide6.QtWidgets import QAbstractItemView, QApplication, QMessageBox

from itmanager.batch_ui import BatchDialog, selected_records, write_export
from itmanager.gui import MainWindow


class Client:
    url = 'http://127.0.0.1:18765'

    def __init__(self):
        self.calls = []

    def get(self, path):
        return []

    def post(self, path, payload):
        self.calls.append((path, payload))
        return {'count': len(payload.get('updates') or payload.get('targets') or [])}


class Runner:
    """Synchronous runner so a test sees the exact payload and can replay failures."""

    def __init__(self):
        self.fail_next = False

    def submit(self, work, done, failed):
        try:
            value = work()
        except Exception as exc:  # pragma: no cover - defensive
            failed(str(exc))
            return
        if self.fail_next:
            self.fail_next = False
            failed('模拟网络失败')
        else:
            done(value)


@pytest.fixture(scope='module')
def qapp():
    return QApplication.instance() or QApplication([])


def records_table(window):
    tables = [t for t in window.page.findChildren(type(window.table([]))) if t.rowCount()]
    assert tables, 'record table with rows expected'
    return tables[0]


def build(qapp):
    window = MainWindow(Client(), autoload=False)
    window.runner = Runner()
    window.records = [
        {'id': 'TASK-1', 'module': 'tasks', 'version': 1, 'archived': 0,
         'data': {'名称': '任务甲', '状态': '未开始', '负责人': '甲'}},
        {'id': 'TASK-2', 'module': 'tasks', 'version': 3, 'archived': 0,
         'data': {'名称': '任务乙', '状态': '未开始', '负责人': '乙'}},
        {'id': 'TASK-3', 'module': 'tasks', 'version': 1, 'archived': 0,
         'data': {'名称': '任务丙', '状态': '已完成', '负责人': '丙'}},
    ]
    window.navigate(window.keys.index('tasks'))
    return window


def test_record_table_allows_multi_selection_and_reports_unique_rows(qapp):
    window = build(qapp)
    table = records_table(window)
    assert table.selectionMode() == QAbstractItemView.ExtendedSelection
    from PySide6.QtCore import QItemSelectionModel as _Selection
    for row in (0, 1):
        table.selectionModel().select(table.model().index(row, 0),
                                      _Selection.Select | _Selection.Rows)
    chosen = selected_records(table)
    assert {row['id'] for row in chosen} == {'TASK-1', 'TASK-2'}
    assert len(selected_records(table)) == 2


def test_field_update_payload_and_idempotent_retry(qapp):
    window = build(qapp)
    dialog = BatchDialog('tasks', window.records[:2], window.client, window.runner, window)
    dialog.mode.setCurrentIndex(0)
    dialog.field.setCurrentText('负责人')
    dialog.value.setCurrentText('统一负责人')
    dialog.show_preview()
    assert '将处理 2 条' in dialog.preview.toPlainText()
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        dialog.apply()
    path, payload = window.client.calls[-1]
    assert path == '/records/batch'
    assert payload['request_id']
    assert [item['version'] for item in payload['updates']] == [1, 3]
    assert all(set(item) == {'id', 'version', 'data'} for item in payload['updates'])
    assert all(item['data']['负责人'] == '统一负责人' for item in payload['updates'])
    first_request = payload['request_id']
    assert dialog.request == {'fingerprint': None, 'id': None}  # success resets the retry key

    # A failed attempt keeps the request id for a byte-identical retry.
    retry = BatchDialog('tasks', window.records[:1], window.client, window.runner, window)
    retry.mode.setCurrentIndex(0)
    retry.field.setCurrentText('负责人')
    retry.value.setCurrentText('失败重试')
    window.runner.fail_next = True
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        retry.apply()
    failed_request = window.client.calls[-1][1]['request_id']
    assert '失败' in retry.message.text()
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        retry.apply()
    assert window.client.calls[-1][1]['request_id'] == failed_request
    # Changing the content must produce a new id rather than replaying the old one.
    retry.value.setCurrentText('内容已改变')
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        retry.apply()
    assert window.client.calls[-1][1]['request_id'] != failed_request


def test_archive_mode_posts_targets_only(qapp):
    window = build(qapp)
    dialog = BatchDialog('tasks', window.records[:2], window.client, window.runner, window)
    dialog.mode.setCurrentIndex(2)
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        dialog.apply()
    path, payload = window.client.calls[-1]
    assert path == '/records/batch-archive'
    assert set(payload) == {'targets', 'request_id'}
    assert payload['targets'] == [{'id': 'TASK-1', 'version': 1}, {'id': 'TASK-2', 'version': 3}]


def test_export_writes_csv_and_xlsx(qapp, tmp_path):
    csv_path = tmp_path / 'out.csv'
    write_export(csv_path, 'tasks', [{'id': 'TASK-1', 'module': 'tasks', 'version': 1,
                                      'data': {'名称': '=危险公式', '下一步': '含,逗号"引号'}}])
    text = csv_path.read_text(encoding='utf-8-sig')
    assert text.splitlines()[0].startswith('记录ID,名称')
    assert "'=危险公式" in text
    assert '"含,逗号""引号"' in text
    xlsx_path = tmp_path / 'out.xlsx'
    write_export(xlsx_path, 'tasks', [{'id': 'TASK-1', 'module': 'tasks', 'version': 1,
                                       'data': {'名称': '演示任务'}}])
    from openpyxl import load_workbook
    sheet = load_workbook(xlsx_path).active
    assert sheet.cell(row=1, column=1).value == '记录ID'
    assert sheet.cell(row=2, column=1).value == 'TASK-1'
    assert sheet.cell(row=2, column=2).value == '演示任务'


def test_batch_entry_without_selection_explains_what_to_do(qapp):
    window = build(qapp)
    table = records_table(window)
    table.clearSelection()
    table.setCurrentCell(-1, -1)
    messages = []
    window.statusBar().showMessage = messages.append
    window.batch_actions('tasks', table)
    assert messages and 'Ctrl' in messages[0]
