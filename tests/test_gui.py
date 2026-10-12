"""Offline client and Qt tests: never calls a real network/model/channel."""
import json
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from itmanager.client import APIClient, ClientError, validate_url, ai_payload, export_data, import_data


class ClientTests(unittest.TestCase):
    def test_transport_policy(self):
        for url in ('http://127.0.0.1:18765', 'http://[::1]:18765', 'http://localhost', 'https://example.com'):
            self.assertEqual(validate_url(url), url)
        for url in ('http://example.com','http://192.168.1.2','https://u:p@example.com','https://example.com?token=x','file:///x','https://example.com\\@localhost'):
            with self.assertRaises(ClientError): validate_url(url)

    def test_token_not_in_repr_or_error(self):
        client = APIClient('http://127.0.0.1', 'secret-token-' * 4)
        self.assertNotIn('secret-token', repr(client))
        with patch('urllib.request.OpenerDirector.open', side_effect=RuntimeError('secret-token-private')):
            with self.assertRaises(ClientError) as context: client.get('/health')
        self.assertNotIn('secret-token', str(context.exception))

    def test_default_content_injected(self):
        record = {'id':'r1','module':'tasks','data':{'名称':'任务'}}
        payload = ai_payload([record], 'p1', 't1')
        self.assertEqual(json.loads(payload['variables']['content']), [record])
        self.assertEqual(payload['variables']['content'], payload['variables']['records'])
        self.assertTrue(payload['confirmed'])
        self.assertFalse(payload['sensitive_confirmed'])

    def test_json_roundtrip(self):
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'export.json'
            payload={'version':1,'records':[{'id':'r','module':'tasks','data':{'名称':'示例'}}]}
            export_data(path,payload)
            self.assertEqual(import_data(path)['records'],payload['records'])

    def test_excel_roundtrip_and_no_formula(self):
        try: from openpyxl import load_workbook
        except ImportError: self.skipTest('openpyxl not installed')
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'export.xlsx'
            payload={'version':1,'records':[{'id':'r','version':1,'module':'tasks','data':{'名称':'=HYPERLINK("x")'}}]}
            export_data(path,payload)
            book=load_workbook(path,data_only=False)
            self.assertEqual(book['任务管理']['D2'].data_type,'s')
            self.assertEqual(len(book.sheetnames),20)
            book.close()
            self.assertEqual(import_data(path)['records'],payload['records'])

    def test_excel_edited_columns_override_json(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'editable.xlsx'
            payload={'records':[{'id':'r','module':'tasks','data':{'名称':'旧名称'}}]}
            export_data(path,payload)
            book=load_workbook(path)
            book['任务管理']['D2']='新名称'
            book.save(path);book.close()
            self.assertEqual(import_data(path)['records'][0]['data']['名称'],'新名称')

    def test_plain_excel_field_mapping(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'plain.xlsx'
            book=Workbook();sheet=book.active;sheet.title='任务管理'
            sheet.append(['名称','负责人','状态']);sheet.append(['任务','张三','未开始'])
            book.save(path);book.close()
            record=import_data(path)['records'][0]
            self.assertEqual(record['module'],'tasks')
            self.assertEqual(record['data']['负责人'],'张三')

try:
    from PySide6.QtCore import QEventLoop, QTimer, QThread
    from PySide6.QtWidgets import QApplication, QComboBox
    from itmanager.gui import MainWindow, RecordDialog, FormDialog, CONFIG_FIELDS, AsyncRunner
    from itmanager.schema import MODULES, STATUS
    QT_AVAILABLE=True
except ImportError:
    QT_AVAILABLE=False


class MockClient:
    url='http://127.0.0.1:18765'
    def __init__(self): self.calls=[]
    def get(self,path):
        self.calls.append(('GET',path))
        if path=='/health':return {'scheduler':False}
        return []
    def post(self,path,body=None):
        self.calls.append(('POST',path,body));return {}


@unittest.skipUnless(QT_AVAILABLE,'PySide6 not installed')
class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.app=QApplication.instance() or QApplication([])

    def test_all_19_forms_and_schema_status(self):
        self.assertEqual(len(MODULES),19)
        for module,(_,fields) in MODULES.items():
            dialog=RecordDialog(module)
            self.assertEqual(set(dialog.widgets),set(fields))
            if '状态' in fields:
                combo=dialog.widgets['状态'][0]
                self.assertIsInstance(combo,QComboBox)
                expected=['目标设定','待评价','待反馈','待复核','已定稿'] if module=='reviews' else STATUS
                self.assertEqual([combo.itemText(i) for i in range(combo.count())],expected)
            dialog.deleteLater()

    def test_final_review_requires_reopening(self):
        record={'id':'review1','data':{'名称':'考核','状态':'已定稿'}}
        dialog=RecordDialog('reviews',record)
        self.assertFalse(dialog.widgets['名称'][0].isEnabled())
        self.assertTrue(dialog.widgets['状态'][0].isEnabled())
        combo=dialog.widgets['状态'][0]
        self.assertEqual([combo.itemText(i) for i in range(combo.count())],['已定稿','待复核'])
        dialog.deleteLater()

    def test_related_selector_uses_id(self):
        records=[{'id':'project1','module':'projects','data':{'名称':'项目一'}}]
        dialog=RecordDialog('tasks',records=records)
        combo=dialog.widgets['项目ID'][0]
        combo.setCurrentIndex(1)
        self.assertEqual(dialog.values()['项目ID'],'project1')
        dialog.deleteLater()

    def test_structured_configuration(self):
        for kind,fields in CONFIG_FIELDS.items():
            dialog=FormDialog(kind,fields)
            self.assertIn('enabled',dialog.widgets)
            self.assertNotIn('JSON',str(dialog.widgets.keys()))
            dialog.deleteLater()

    def test_window_navigation_without_network(self):
        client=MockClient(); window=MainWindow(client,autoload=False)
        for module in MODULES:
            window.nav.setCurrentRow(window.keys.index(module))
            self.assertEqual(window.current,module)
        self.assertEqual(client.calls,[])
        window.close();window.deleteLater()

    def test_async_callback_is_gui_thread(self):
        # Assert the thread affinity of worker vs callback. A long-but-generous timeout
        # keeps this honest: the full suite creates many windows, which slows the Qt
        # event loop, so a 3s window made this assertion flaky without indicating any
        # real defect (a single-window process schedules the job in milliseconds).
        runner=AsyncRunner(); loop=QEventLoop(); observed=[]
        def work():return QThread.currentThread()==self.app.thread()
        def done(value):
            observed.append((value,QThread.currentThread()==self.app.thread()));loop.quit()
        runner.submit(work,done,lambda error:loop.quit())
        QTimer.singleShot(20000,loop.quit);loop.exec()
        self.assertEqual(observed,[(False,True)])
        self.assertFalse(runner.jobs)


if __name__=='__main__':unittest.main()
